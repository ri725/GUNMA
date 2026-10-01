from __future__ import annotations

import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "工場の質問箱_質問登録.xlsx"
EMPTY = {"", "（空）", None}
TOKYO = ZoneInfo("Asia/Tokyo")
_lock = threading.Lock()


@dataclass
class Item:
    number: int
    category: str
    kind: str
    question: str
    aliases: str
    answer: str
    after: str
    owner: str
    state: str

    @property
    def active(self) -> bool:
        return self.state != "止める"


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def menu_items(category: str) -> list[str]:
    with _lock:
        wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
        try:
            sheet = wb["分類"]
            headers = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            if category not in headers:
                return []
            col = headers.index(category) + 1
            items = []
            for (value,) in sheet.iter_rows(min_row=2, min_col=col, max_col=col, values_only=True):
                label = _text(value)
                if label in EMPTY:
                    continue
                items.append(label)
            return items
        finally:
            wb.close()


def questions() -> list[Item]:
    with _lock:
        wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
        try:
            sheet = wb["質問箱"]
            rows = []
            for index, row in enumerate(sheet.iter_rows(min_row=2, max_col=10, values_only=True), start=2):
                item = Item(
                    number=index - 1,
                    category=_text(row[1]),
                    kind=_text(row[2]),
                    question=_text(row[3]),
                    aliases=_text(row[4]),
                    answer=_text(row[5]),
                    after=_text(row[6]),
                    owner=_text(row[7]),
                    state=_text(row[8]),
                )
                if not any([item.category, item.kind, item.question, item.answer]):
                    continue
                rows.append(item)
            return rows
        finally:
            wb.close()


def answers_for(category: str, kind: str) -> list[Item]:
    return [
        item
        for item in questions()
        if item.active and item.category == category and item.kind == kind and item.answer
    ]


def log_incoming(
    text: str,
    category: str = "",
    kind: str = "",
    result: str = "本人待ち",
    used_number: str = "",
) -> None:
    with _lock:
        wb = load_workbook(WORKBOOK)
        try:
            sheet = wb["届いた質問"]
            target = None
            for row in range(2, sheet.max_row + 1):
                if not _text(sheet.cell(row, 4).value):
                    target = row
                    break
            if target is None:
                target = sheet.max_row + 1
            now = datetime.now(TOKYO).strftime("%Y/%m/%d %H:%M")
            values = [now, category, kind, text, result, used_number, "", "まだ" if result != "返した" else ""]
            for col, value in enumerate(values, start=1):
                sheet.cell(target, col, value)
            wb.save(WORKBOOK)
        finally:
            wb.close()
